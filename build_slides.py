"""Builds P30_slides.pptx from the project's saved results. A one-off
report-generation script (not part of the src/ analysis pipeline) --
mirrors REPORT.md's structure and numbers, both drawn from the same
verified results/tables/*.csv and research_log.md sources.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TITLE_COLOR = RGBColor(0x1B, 0x3A, 0x5C)
ACCENT_COLOR = RGBColor(0xC0, 0x39, 0x2B)
TEXT_COLOR = RGBColor(0x22, 0x22, 0x22)

FIG = "results/figures"


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(12), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(["P30: Koopman/EDMD Forecasting of MRSA Spread", "on a Hospital Contact Network"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(34)
        p.font.bold = True
        p.font.color.rgb = TITLE_COLOR

    sub = slide.shapes.add_textbox(Inches(0.7), Inches(3.9), Inches(12), Inches(1.5))
    tf2 = sub.text_frame
    tf2.word_wrap = True
    for text, size in [
        ("CMDS Course Project · M.Tech Data Science · Amrita Vishwa Vidyapeetham", 18),
        ("Topic P30 · First review: October 6, 2026", 16),
    ]:
        p = tf2.add_paragraph() if tf2.paragraphs[0].text else tf2.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = TEXT_COLOR
    return slide


def add_section_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = TITLE_COLOR
    box = slide.shapes.add_textbox(Inches(0.7), Inches(3.0), Inches(12), Inches(2))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(title.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    return slide


def add_bullet_slide(prs, title, bullets, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.9))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(26)
    tp.font.bold = True
    tp.font.color.rgb = TITLE_COLOR

    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12.1), Inches(5.6))
    tf = body.text_frame
    tf.word_wrap = True
    first = True
    for item in bullets:
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = ("• " if level == 0 else "   - ") + text
        p.font.size = Pt(20 if level == 0 else 17)
        p.font.color.rgb = ACCENT_COLOR if level == 0 and text.isupper() else TEXT_COLOR
        p.space_after = Pt(10)
    if note:
        note_box = slide.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(12.1), Inches(0.5))
        p = note_box.text_frame.paragraphs[0]
        p.text = note
        p.font.size = Pt(12)
        p.font.italic = True
        p.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    return slide


def add_image_slide(prs, title, image_path, caption=None, bullets=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.8))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(26)
    tp.font.bold = True
    tp.font.color.rgb = TITLE_COLOR

    left_width = Inches(7.6) if bullets else Inches(11.5)
    pic = slide.shapes.add_picture(image_path, Inches(0.6), Inches(1.25), width=left_width)
    # cap height so it never overflows the slide
    max_h = Inches(5.6)
    if pic.height > max_h:
        ratio = max_h / pic.height
        pic.height = max_h
        pic.width = Emu(int(pic.width * ratio))

    if bullets:
        body = slide.shapes.add_textbox(Inches(8.5), Inches(1.4), Inches(4.3), Inches(5.4))
        tf = body.text_frame
        tf.word_wrap = True
        first = True
        for text in bullets:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = "• " + text
            p.font.size = Pt(15)
            p.font.color.rgb = TEXT_COLOR
            p.space_after = Pt(8)

    if caption:
        cap_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(11.5), Inches(0.4))
        p = cap_box.text_frame.paragraphs[0]
        p.text = caption
        p.font.size = Pt(12)
        p.font.italic = True
        p.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    return slide


def add_table_slide(prs, title, headers, rows, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.8))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(26)
    tp.font.bold = True
    tp.font.color.rgb = TITLE_COLOR

    n_rows, n_cols = len(rows) + 1, len(headers)
    table_shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.5 * n_rows))
    table = table_shape.table
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(15)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val)
            cell.text_frame.paragraphs[0].font.size = Pt(14)

    if note:
        note_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.5 + 0.5 * n_rows), Inches(12.1), Inches(1.0))
        p = note_box.text_frame.paragraphs[0]
        p.text = note
        p.font.size = Pt(14)
        p.font.italic = True
        p.font.color.rgb = TEXT_COLOR
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)

    add_bullet_slide(
        prs, "Motivation",
        [
            "Hospital-acquired MRSA spreads through person-to-person contact",
            "Infection-control resources (isolation, cohorting, screening) are finite",
            "Which ward, if any, is disproportionately responsible for spread?",
            "Would intervening there actually help more than a blanket or random policy?",
        ],
    )

    add_bullet_slide(
        prs, "Research Questions",
        [
            "RQ1: Does lifting ward-level MRSA prevalence into a Koopman/EDMD",
            ("framework forecast outbreak trajectories better than simple baselines?", 1),
            "RQ2: Do the resulting Koopman eigenmodes identify “superspreader”",
            ("wards better than plain network centrality?", 1),
            "Answered honestly against a SIMULATED GROUND TRUTH, not just in-sample fit",
        ],
    )

    add_bullet_slide(
        prs, "Data: the I-Bird Hospital Contact Network",
        [
            "200-bed long-term/rehabilitation hospital, Berck-sur-Mer, France, 2009",
            "Obadia et al. 2015 (PLOS Comput Biol); Duval et al. 2018 (Sci Rep)",
            "admission.csv: 795 people (452 patients, 343 staff), 6 ward groups",
            "mat.day.csv: 124,924 contact records, 589 people, 19,974 unique pairs",
            "microbio.csv: 6,728 MRSA swabs, 295 people MRSA-positive at least once",
            "P1 data audit: all 19 reference numbers independently re-verified exactly",
        ],
    )

    add_image_slide(
        prs, "The Contact Network, Colored by Ward",
        f"{FIG}/p2_network_by_ward_all.png",
        bullets=[
            "Clear within-ward clustering",
            "“Other” (staff, brown) bridges wards",
            "90.3% of contact-seconds are within-ward",
            "Ward contact matrix W built from raw bidirectional contact records",
        ],
        caption="Figure: 589-person contact network, node color = ward group.",
    )

    add_bullet_slide(
        prs, "Pipeline Overview",
        [
            "P1 Data audit  →  P2 Contact network + ward matrix W",
            "P3 Real weekly prevalence  →  P4 Individual-level SIS simulation + calibration",
            "P5 EDMD dictionaries (D1/D2/D3)  →  P6 Ridge-regularized EDMD fit",
            "P7 Eigenmode analysis  →  P8 Real-data forecasting validation",
            "P9 Superspreader risk score vs. centrality  →  P10 Intervention simulation",
            "P11 Robustness experiments E1–E8",
            "179 automated tests throughout; every design decision logged and verified",
        ],
    )

    add_bullet_slide(
        prs, "Individual-Level SIS Simulation (P4)",
        [
            "Daily stochastic SIS on the 589-person contact network",
            "P(transmit) = 1 − exp(−β·contact hours with infected neighbors)",
            "Verified analytically: matched a closed-form solution over 20,000 replicates",
            "Calibrated to real weekly prevalence (weeks 0–11), 320-point grid search",
            "FINDING: calibration sits on a wide, weakly-identified loss ridge",
            ("— only ~16 real weekly snapshots to calibrate against", 1),
            "300 simulated training trajectories generated for EDMD fitting",
        ],
        note="Best fit: β=6.8×10⁻⁴/hr, γ=1.7×10⁻³/day — 16 of 320 grid points within 10% of best loss",
    )

    add_bullet_slide(
        prs, "Three EDMD Dictionaries (P5)",
        [
            "D1 (linear): [1, x] — 7 features",
            "D2 (degree-2 polynomial): [1, x, x², cross terms] — 28 features",
            "D3 (contact-weighted): [1, x, x², Wx, x⊙(Wx)] — 25 features",
            "D3 is the one dictionary meant to inject network structure",
        ],
    )

    add_bullet_slide(
        prs, "Major Finding: D3 Adds No Function Space Beyond D2",
        [
            "PROVEN ANALYTICALLY, before building anything on top of it:",
            ("Wx is a linear function of x → spans exactly the same subspace as x", 1),
            ("x·(Wx) is a fixed linear combination of D2's own square/cross terms", 1),
            ("→ D3's effective rank ceiling: 19 of its 25 raw features", 1),
            "CONFIRMED EMPIRICALLY: fitted D3 operator has rank(K) = 19 exactly",
            "D1 (rank 7/7) and D2 (rank 28/28) are both full rank",
            "Consequence: D3 tests whether a contact-structured PRIOR helps —",
            ("never “injects network information,” since it structurally cannot", 1),
        ],
    )

    add_bullet_slide(
        prs, "EDMD Fit Results (P6)",
        [
            "Ridge regression, written from scratch in NumPy (no external Koopman library)",
            "Constant feature excluded from scaling AND the ridge penalty",
            "λ selected on validation split, mean RMSE over forecast horizons 1–4",
            "HONEST RESULT: D1, D2, D3 perform almost identically",
            ("Test h=1 RMSE: D1=0.00748, D2=0.00758, D3=0.00747", 1),
            ("All beat persistence by 12–13% at h=1, growing to 25–27% by h=4", 1),
            "D3 edges out D2 very slightly, despite having less effective capacity",
        ],
    )

    add_image_slide(
        prs, "Eigenmode Analysis (P7): a Disagreement Worth Tracking",
        f"{FIG}/p7_eigenvalue_spectrum_D3.png",
        bullets=[
            "Dominant (most persistent) mode's top ward:",
            "D1 says Menard 1 (45.2% share)",
            "D2 and D3 agree: Sorrel 1 (~42% share)",
            "D3's 6 spurious null-space eigenvalues",
            "(from the proven rank deficiency) sit",
            "near the origin, cleanly separated",
            "from the 19 real modes near |λ|=1",
            "This disagreement is NOT resolved from",
            "eigenmodes alone — tested next against",
            "simulated ground truth",
        ],
        caption="Figure: D3's eigenvalue spectrum — real modes (blue) vs. spurious null-space modes (red).",
    )

    add_bullet_slide(
        prs, "Real-Data Forecasting Validation (P8)",
        [
            "First time the fitted K touches real data (deliberately held back until now)",
            "New baselines built: network-exposure linear model, ward-level SIS model",
            "HONEST RESULT on the holdout period (weeks 12–16):",
            ("D1/D2/D3 statistically indistinguishable from persistence at h=1", 1),
            ("historical_mean overtakes everything by h=3–4 (skill ≈0.20)", 1),
            "Plausible cause: K trained on simulated trajectories sitting on P4's",
            ("weakly-identified calibration ridge — may not transfer a strong edge", 1),
        ],
    )

    add_section_slide(prs, "The Central Question:\nKoopman Eigenmodes vs. Simple Centrality")

    add_image_slide(
        prs, "Simulated Superspreader Ground Truth (P9)",
        f"{FIG}/p9_ground_truth_bar.png",
        bullets=[
            "Ground truth: cut each ward's",
            "contacts 50%, measure the drop",
            "in colonized person-days",
            "(100 replicates, common random",
            "numbers, paired comparison)",
            "",
            "Ranking: Menard 1 > Sorrel 1 >",
            "Menard 2 > Sorrel 2 > Sorrel 0 >",
            "Other",
        ],
        caption="Figure: drop in colonized person-days per ward, error bars = per-replicate std.",
    )

    add_table_slide(
        prs, "Risk Score vs. Ground Truth: Spearman ρ",
        ["Candidate", "ρ", "p-value"],
        [
            ["degree_centrality", "0.943", "0.0048"],
            ["eigenvector_centrality", "0.714", "0.111"],
            ["koopman_D2", "0.543", "0.266"],
            ["koopman_D3", "0.429", "0.397"],
            ["betweenness_centrality", "0.309", "0.552"],
            ["koopman_D1", "0.257", "0.623"],
        ],
        note="Simple degree centrality (a ward's raw contact volume, no Koopman machinery) predicts the ground\ntruth far better than any Koopman eigenmode score — the only candidate surviving Bonferroni correction.",
    )

    add_image_slide(
        prs, "The Same Pattern, Visually",
        f"{FIG}/p9_rank_heatmap.png",
        bullets=[
            "ground_truth, degree_centrality,",
            "and eigenvector_centrality",
            "columns look nearly identical",
            "",
            "The three koopman columns",
            "look comparatively scrambled",
            "",
            "Nuance: koopman_D1 DOES pick",
            "the #1 ward correctly (Menard 1)",
            "— but has the worst OVERALL",
            "rank correlation of the three",
        ],
        caption="Figure: ward ranks (1=highest) across all candidates and the ground truth.",
    )

    add_image_slide(
        prs, "Intervention Simulation (P10): Targeting Matters",
        f"{FIG}/p10_drop_vs_budget.png",
        bullets=[
            "Equal-budget framework:",
            "budget = k × reduction fraction",
            "",
            "Every TARGETED strategy beats",
            "BOTH untargeted strategies",
            "(random, whole-hospital) at",
            "every matching budget",
            "",
            "At k=2: koopman_D2/D3 slightly",
            "beats degree-targeting — a higher",
            "overall correlation doesn't",
            "guarantee the better k=2 pick",
        ],
        caption="Figure: drop in colonized person-days vs. budget, by strategy.",
    )

    add_table_slide(
        prs, "Robustness: Experiments E1–E8 (P11)",
        ["#", "Probes", "Result"],
        [
            ["E1", "Seed sensitivity", "ρ=0.943"],
            ["E2", "λ sensitivity (0.1×/10×)", "max ratio dev. 3.3%"],
            ["E3", "Replicate count (30/300)", "ρ=1.000"],
            ["E4", "Calibration ridge-point", "ρ=1.000"],
            ["E5", "W's time window", "ρ=1.000"],
            ["E6", "Patients-only population", "ρ≈1.000"],
            ["E7", "Missing-data handling", "ρ=1.000 (model ranking)"],
            ["E8", "D3 structural redundancy", "ratio=0.998, same dominant ward"],
        ],
        note="Every major conclusion is stable across seed, replicate count, calibration point, network\ntime-window, population, missing-data handling, and λ choice.",
    )

    add_bullet_slide(
        prs, "Discussion",
        [
            "The Koopman/EDMD machinery works and was built/verified rigorously",
            "(exact recovery on synthetic systems, closed-form agreement, 179 tests)",
            "On THIS dataset, it does not outperform simpler alternatives — reported",
            "honestly, with three concrete, verifiable reasons:",
            ("(a) SIS calibration sits on a weakly-identified ridge (~16 real snapshots)", 1),
            ("(b) D3, the “network-aware” dictionary, is PROVABLY no richer than D2", 1),
            ("(c) raw ward contact volume is a strong, nearly-free predictor", 1),
            "Targeting still matters enormously, independent of which score is used —",
            "a practically actionable finding that survives every robustness check",
        ],
    )

    add_bullet_slide(
        prs, "Limitations",
        [
            "About 16 real weekly snapshots only — root cause of the calibration ridge",
            "Time-invariant K vs. a time-varying contact network",
            "Only 6 ward groups — rank statistics are inherently weak (n=6)",
            "No admission/discharge dates, single hospital, single 4-month window",
            "No transfer records",
        ],
    )

    add_bullet_slide(
        prs, "Conclusion & Future Work",
        [
            "Built a complete, rigorously-tested pipeline: real network → calibrated",
            "SIS simulator → EDMD → eigenmodes → validated against simulated ground truth",
            "Honest answer: network-aware Koopman eigenmodes do not beat simple",
            "degree centrality for THIS dataset — traced to specific, understood causes",
            "“Some form of targeting beats none” is robust and practically useful",
            "FUTURE WORK:",
            ("Genuinely time-varying, network-weighted dictionary (week-specific W_t)", 1),
            ("Larger / multi-site real dataset to better identify the SIS calibration", 1),
        ],
    )

    add_bullet_slide(
        prs, "References",
        [
            "Obadia, T., Silhol, R., Opatowski, L., Temime, L., Legrand, J.,",
            "Thiébaut, A.C.M., et al. (2015). Detailed Contact Data and the",
            "Dissemination of Staphylococcus aureus in Hospitals.",
            "PLOS Computational Biology, 11(3):e1004170.",
            "",
            "Duval, A., Obadia, T., Martinet, L., Boëlle, P.Y., Fleury, E.,",
            "Guillemot, D., et al. (2018). Measuring dynamic social contacts",
            "in a rehabilitation hospital: effect of wards, patient and staff",
            "characteristics. Scientific Reports, 8:1686.",
        ],
    )

    prs.save("P30_slides.pptx")
    print(f"Saved P30_slides.pptx with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()
